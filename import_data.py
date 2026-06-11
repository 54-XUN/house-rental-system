import os
import sys
import gc

sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'backend'))

from app import 创建应用
from models import db
from models.house import 房源模型
from models.customer import 客户模型
from models.contract import 合同模型

app = 创建应用()

def 解析日期(值):
    if not 值:
        return None
    try:
        if hasattr(值, 'strftime'):
            return 值.strftime('%Y-%m-%d')
        return str(值)[:10] if 值 else None
    except:
        return None

def 安全转整数(值, 默认=None):
    if 值 is None:
        return 默认值
    try:
        return int(值)
    except:
        return 默认值

def 安全转浮点(值, 默认=None):
    if 值 is None:
        return 默认
    try:
        return float(值)
    except:
        return 默认

with app.app_context():
    excel路径 = os.path.join(os.path.dirname(__file__), '房屋出租管理系统.xlsx')
    
    print("🧹 清理旧数据...")
    db.session.query(合同模型).delete()
    db.session.query(房源模型).delete()
    db.session.query(客户模型).delete()
    db.session.commit()
    
    import openpyxl
    工作簿 = openpyxl.load_workbook(excel路径, read_only=True, data_only=True)
    
    # ========== 导入房源 ==========
    if '房源明细' in 工作簿.sheetnames:
        print("🏠 导入房源...")
        工作表 = 工作簿['房源明细']
        已用编号 = set()
        计数 = 0
        
        for 行 in 工作表.iter_rows(min_row=4, max_col=15, values_only=True):
            if not any(行[:2]):
                continue
            
            编号 = str(行[1]).strip() if len(行) > 1 and 行[1] else None
            if not 编号 or not 编号.startswith('FY'):
                continue
            if 编号 in 已用编号:
                continue
            已用编号.add(编号)
            
            房源 = 房源模型(
                house_code=编号,
                add_date=解析日期(行[0]),
                community=str(行[2]).strip() if len(行) > 2 and 行[2] else None,
                address=str(行[3]).strip() if len(行) > 3 and 行[3] else None,
                floor=安全转整数(行[4]),
                room=安全转整数(行[5]),
                hall=安全转整数(行[6]),
                area=安全转浮点(行[7]),
                rent=安全转浮点(行[8]),
                tags=str(行[9]).replace('|', ',') if len(行) > 9 and 行[9] else None,
                status=str(行[12]).strip() if len(行) > 12 and 行[12] else '空闲'
            )
            db.session.add(房源)
            计数 += 1
            if 计数 % 20 == 0:
                db.session.commit()
        
        db.session.commit()
        print(f"  ✅ {计数} 条")
    
    # ========== 导入客户 ==========
    if '客户管理' in 工作簿.sheetnames:
        print("👥 导入客户...")
        工作表 = 工作簿['客户管理']
        已用编号 = set()
        计数 = 0
        
        for 行 in 工作表.iter_rows(min_row=4, max_col=10, values_only=True):
            if not any(行[:2]):
                continue
            
            编号 = str(行[1]).strip() if len(行) > 1 and 行[1] else None
            if not 编号 or not 编号.startswith('F'):
                continue
            if 编号 in 已用编号:
                continue
            已用编号.add(编号)
            
            客户 = 客户模型(
                customer_code=编号,
                add_date=解析日期(行[0]),
                name=str(行[2]).strip() if len(行) > 2 and 行[2] else None,
                id_card=str(行[3]).strip() if len(行) > 3 and 行[3] else None,
                phone=str(行[4]).strip() if len(行) > 4 and 行[4] else None,
                wechat=str(行[5]).strip() if len(行) > 5 and 行[5] else None,
                tags=str(行[6]).replace('|', ',') if len(行) > 6 and 行[6] else None,
                status=str(行[7]).strip() if len(行) > 7 and 行[7] else '跟进中'
            )
            db.session.add(客户)
            计数 += 1
            if 计数 % 20 == 0:
                db.session.commit()
        
        db.session.commit()
        print(f"  ✅ {计数} 条")
    
    # ========== 导入合同 ==========
    if '交易流水' in 工作簿.sheetnames:
        print("📝 导入合同...")
        工作表 = 工作簿['交易流水']
        所有房源 = {h.house_code: h for h in 房源模型.query.all()}
        计数 = 0
        跳过 = 0
        
        for 行 in 工作表.iter_rows(min_row=4, max_col=12, values_only=True):
            if not any(行[:2]):
                continue
            
            房源编号 = str(行[7]).strip() if len(行) > 7 and 行[7] else None
            if not 房源编号 or 房源编号 not in 所有房源:
                跳过 += 1
                continue
            
            合同编号 = str(行[1]).strip() if len(行) > 1 and 行[1] else None
            客户编号 = str(行[9]).strip() if len(行) > 9 and 行[9] else None
            
            合同 = 合同模型(
                contract_code=合同编号,
                sign_date=解析日期(行[0]),
                house_code=房源编号,
                customer_code=客户编号,
                months=安全转整数(行[2], 0),
                start_date=解析日期(行[3]),
                end_date=解析日期(行[4]),
                deposit=安全转浮点(行[5]),
                monthly_rent=安全转浮点(行[8]),
                total_amount=安全转浮点(行[6])
            )
            db.session.add(合同)
            
            关联房源 = 所有房源[房源编号]
            关联房源.contract_code = 合同编号
            关联房源.expire_date = 解析日期(行[4])
            关联房源.status = '已租'
            
            计数 += 1
            if 计数 % 10 == 0:
                db.session.commit()
        
        db.session.commit()
        print(f"  ✅ {计数} 条 (跳过 {跳过} 条)")
    
    工作簿.close()
    gc.collect()
    
    print("\n" + "="*40)
    print("📊 导入完成!")
    print(f"  🏠 房源: {房源模型.query.count()} 条")
    print(f"  👥 客户: {客户模型.query.count()} 条")
    print(f"  📝 合同: {合同模型.query.count()} 条")
